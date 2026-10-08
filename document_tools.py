"""Document extraction and rule-based checks for thesis files."""

from __future__ import annotations

import re
import zipfile
from collections import Counter
from io import BytesIO
from typing import Any

from docx import Document
import pdfplumber

MAX_UPLOAD_BYTES = 4_400_000
MAX_DOCX_UNPACKED_BYTES = 32 * 1024 * 1024


def _check(name: str, status: str, detail: str) -> dict[str, str]:
    return {"name": name, "status": status, "detail": detail}


def _near(value: float, expected: float, tolerance: float = 0.08) -> bool:
    return abs(value - expected) <= tolerance


def _citation_audit(text: str, requested_style: str) -> dict[str, Any]:
    reference_heading = re.search(
        r"(?im)^\s*(?:บรรณานุกรม|เอกสารอ้างอิง|รายการเอกสารอ้างอิง|references|bibliography)\s*$",
        text,
    )
    body = text[:reference_heading.start()] if reference_heading else text
    references = text[reference_heading.end():] if reference_heading else ""
    ieee_groups = re.findall(r"\[([0-9][0-9,\s\-]*)\]", body)
    apa_citations = re.findall(r"\(([^()\n]{1,120}?(?:19|20)\d{2}[a-z]?(?:[^()\n]{0,40})?)\)", body)

    style = requested_style.lower()
    if style == "auto":
        style = "ieee" if ieee_groups and len(ieee_groups) >= len(apa_citations) else "apa7"

    result: dict[str, Any] = {
        "style": style,
        "reference_section_found": bool(reference_heading),
        "in_text_citation_count": len(ieee_groups) if style == "ieee" else len(apa_citations),
        "reference_entry_count": 0,
        "issues": [],
        "note": "รูปแบบ APA/IEEE ตรวจด้วยกฎเบื้องต้น ควรตรวจรายการสุดท้ายด้วยคน",
    }
    if not reference_heading:
        result["issues"].append("ไม่พบหัวข้อบรรณานุกรมหรือเอกสารอ้างอิง")
        return result

    if style == "ieee":
        listed = {int(number) for number in re.findall(r"(?m)^\s*\[(\d+)\]\s*\S", references)}
        cited: set[int] = set()
        for group in ieee_groups:
            for part in re.split(r"[,\s]+", group.strip()):
                if not part:
                    continue
                if "-" in part:
                    try:
                        first, last = (int(value) for value in part.split("-", 1))
                        if last - first <= 100:
                            cited.update(range(first, last + 1))
                    except ValueError:
                        continue
                elif part.isdigit():
                    cited.add(int(part))
        result["reference_entry_count"] = len(listed)
        result["cited_numbers_missing_from_references"] = sorted(cited - listed)
        result["references_never_cited"] = sorted(listed - cited)
        if cited - listed:
            result["issues"].append("มีหมายเลขอ้างอิงในเนื้อหาที่ไม่พบในรายการท้ายเล่ม")
        if listed - cited:
            result["issues"].append("มีรายการอ้างอิงท้ายเล่มที่ไม่พบหมายเลขอ้างถึงในเนื้อหา")
    else:
        entries = [line.strip() for line in references.splitlines() if line.strip()]
        years_in_entries = set(re.findall(r"\b(?:19|20)\d{2}[a-z]?\b", references))
        citation_years = set(re.findall(r"\b(?:19|20)\d{2}[a-z]?\b", " ".join(apa_citations)))
        result["reference_entry_count"] = len(entries)
        result["citation_years_missing_from_references"] = sorted(citation_years - years_in_entries)
        if citation_years - years_in_entries:
            result["issues"].append("พบปีในอ้างอิงในเนื้อหาที่ไม่พบปีเดียวกันในรายการท้ายเล่ม")
        if not apa_citations:
            result["issues"].append("ไม่พบรูปแบบอ้างอิงนาม-ปีที่ตรวจจับได้ อาจเป็นเพราะใช้รูปแบบอื่น")
    return result


def _outline_audit(text: str) -> dict[str, Any]:
    chapter_patterns = {
        "บทที่ 1 บทนำ": r"(?:บทที่\s*1\b|chapter\s*1\b|บทนำ|introduction)",
        "บทที่ 2 เอกสารและงานวิจัยที่เกี่ยวข้อง": r"(?:บทที่\s*2\b|chapter\s*2\b|เอกสารและงานวิจัยที่เกี่ยวข้อง|literature review)",
        "บทที่ 3 วิธีดำเนินการวิจัย": r"(?:บทที่\s*3\b|chapter\s*3\b|วิธีดำเนินการวิจัย|ระเบียบวิธีวิจัย|methodology)",
        "บทที่ 4 ผลการวิจัย": r"(?:บทที่\s*4\b|chapter\s*4\b|ผลการวิจัย|results)",
        "บทที่ 5 สรุป อภิปรายผลและข้อเสนอแนะ": r"(?:บทที่\s*5\b|chapter\s*5\b|สรุป.*อภิปรายผล|conclusions?.*discussion)",
    }
    chapters = [
        {"name": title, "status": "found" if re.search(pattern, text, re.I) else "missing"}
        for title, pattern in chapter_patterns.items()
    ]
    elements = [
        ("ความเป็นมาและความสำคัญของปัญหา", r"ความเป็นมา|ความสำคัญของปัญหา|background.*problem"),
        ("วัตถุประสงค์", r"วัตถุประสงค์|objectives?"),
        ("ขอบเขตการวิจัย", r"ขอบเขตของการวิจัย|ขอบเขตการวิจัย|scope of (?:the )?research"),
        ("สมมติฐาน (ถ้ามี)", r"สมมติฐาน|hypotheses?"),
        ("เอกสารและงานวิจัยที่เกี่ยวข้อง", r"เอกสารและงานวิจัยที่เกี่ยวข้อง|literature review"),
        ("วิธีดำเนินการวิจัยและเครื่องมือ", r"วิธีดำเนินการวิจัย|เครื่องมือที่ใช้ในการวิจัย|research methodology"),
        ("การเก็บรวบรวมและวิเคราะห์ข้อมูล", r"เก็บรวบรวมข้อมูล|วิเคราะห์ข้อมูล|data collection|data analysis"),
        ("ผลการวิจัย", r"ผลการวิจัย|results"),
        ("สรุป อภิปรายผล และข้อเสนอแนะ", r"สรุป.*อภิปรายผล|ข้อเสนอแนะ|conclusions?.*discussion"),
    ]
    checks = [
        {"name": title, "status": "found" if re.search(pattern, text, re.I | re.S) else "review", "detail": "พบข้อความที่เกี่ยวข้อง" if re.search(pattern, text, re.I | re.S) else "ไม่พบคำสำคัญนี้โดยอัตโนมัติ โปรดตรวจสารบัญและเนื้อหา"}
        for title, pattern in elements
    ]
    return {"chapters": chapters, "elements": checks}


def _analyze_docx(raw: bytes) -> tuple[str, list[dict[str, str]], dict[str, Any]]:
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_DOCX_UNPACKED_BYTES:
            raise ValueError("ไฟล์ Word มีขนาดหลังแตกไฟล์ใหญ่เกินกำหนด")
    document = Document(BytesIO(raw))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_text = [cell.text.strip() for table in document.tables for row in table.rows for cell in row.cells if cell.text.strip()]
    header_footer_text = []
    page_fields = False
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer):
            header_footer_text.extend(p.text.strip() for p in part.paragraphs if p.text.strip())
            if re.search(r"\bPAGE\b", part._element.xml, re.I):
                page_fields = True
    text = "\n".join(paragraphs + table_text + header_footer_text)

    checks = []
    for index, section in enumerate(document.sections, start=1):
        margins = {
            "บน": section.top_margin.inches if section.top_margin else None,
            "ล่าง": section.bottom_margin.inches if section.bottom_margin else None,
            "ซ้าย": section.left_margin.inches if section.left_margin else None,
            "ขวา": section.right_margin.inches if section.right_margin else None,
        }
        expected = {"บน": 1.5, "ล่าง": 1.0, "ซ้าย": 1.5, "ขวา": 1.0}
        mismatched = [f"{side} {value:.2f} นิ้ว" for side, value in margins.items() if value is None or not _near(value, expected[side])]
        checks.append(_check(
            f"ระยะขอบ (ส่วนที่ {index})",
            "pass" if not mismatched else "review",
            "ตรงกับคู่มือ (ขอบบน/ซ้าย 1.5 นิ้ว, ล่าง/ขวา 1 นิ้ว)" if not mismatched else "ค่าที่พบ: " + (", ".join(mismatched) or "ไม่สามารถอ่านได้") + "; หน้าเริ่มบทอาจกำหนดขอบบน 2 นิ้ว",
        ))
        width = section.page_width.inches if section.page_width else 0
        height = section.page_height.inches if section.page_height else 0
        is_a4 = _near(width, 8.27, 0.12) and _near(height, 11.69, 0.12)
        checks.append(_check(f"กระดาษ A4 (ส่วนที่ {index})", "pass" if is_a4 else "review", f"ขนาด {width:.2f} × {height:.2f} นิ้ว"))

    font_counts: Counter[str] = Counter()
    sizes: Counter[float] = Counter()
    for paragraph in document.paragraphs:
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            continue
        for run in paragraph.runs:
            if not run.text.strip():
                continue
            if run.font.name:
                font_counts[run.font.name] += len(run.text.strip())
            if run.font.size:
                sizes[round(run.font.size.pt, 1)] += len(run.text.strip())
    if font_counts:
        common_font, _ = font_counts.most_common(1)[0]
        allowed_fonts = {"angsana upc", "angsanaUPC".lower(), "browallia upc", "browalliaUPC".lower(), "dillenia upc", "dilleniaUPC".lower(), "eucrosia upc", "eucrosiaUPC".lower(), "cordia upc", "cordiaUPC".lower(), "th sarabunpsk", "times new roman"}
        status = "pass" if common_font.strip().lower() in allowed_fonts else "review"
        checks.append(_check("แบบอักษรหลัก", status, f"ตรวจพบ {common_font}; รายชื่อที่คู่มือระบุรวม AngsanaUPC, BrowalliaUPC, DilleniaUPC, EucrosiaUPC, CordiaUPC, TH SarabunPSK และ Times New Roman"))
    else:
        checks.append(_check("แบบอักษรหลัก", "review", "ไฟล์ไม่ได้กำหนดชื่อฟอนต์ไว้ในข้อความโดยตรง จึงตรวจเทียบสไตล์ที่สืบทอดไม่ได้ครบ"))
    if sizes:
        common_size, _ = sizes.most_common(1)[0]
        checks.append(_check("ขนาดตัวอักษรเนื้อหา", "pass" if min(abs(common_size - 16), abs(common_size - 12)) <= 1 else "review", f"ขนาดที่พบบ่อย {common_size:g} พอยต์; คู่มือระบุ 16 พอยต์สำหรับภาษาไทย และ 12 พอยต์สำหรับภาษาอังกฤษ"))
    else:
        checks.append(_check("ขนาดตัวอักษรเนื้อหา", "review", "ไม่พบขนาดตัวอักษรที่กำหนดโดยตรงในย่อหน้า"))
    checks.append(_check("เลขหน้า", "pass" if page_fields else "review", "พบฟิลด์เลขหน้า PAGE ในหัวกระดาษ/ท้ายกระดาษ" if page_fields else "ไม่พบฟิลด์เลขหน้า PAGE ในหัวกระดาษ/ท้ายกระดาษ"))
    approval_found = bool(re.search(r"ใบรับรองวิทยานิพนธ์|คณะกรรมการสอบวิทยานิพนธ์", text, re.I))
    checks.append(_check("หน้าอนุมัติ/ใบรับรอง", "pass" if approval_found else "review", "พบคำว่าใบรับรองหรือคณะกรรมการสอบในเอกสาร" if approval_found else "ไม่พบข้อความที่บ่งชี้หน้าใบรับรอง โปรดตรวจด้วยตนเอง"))
    return text, checks, {"page_count": None, "page_count_note": "จำนวนหน้า Word ขึ้นกับโปรแกรมและเครื่องพิมพ์; ไม่สามารถนับได้แน่นอนจากไฟล์ต้นฉบับ"}


def _analyze_pdf(raw: bytes) -> tuple[str, list[dict[str, str]], dict[str, Any]]:
    texts = []
    page_sizes = []
    edge_number_pages = 0
    approval_found = False
    font_sizes: Counter[float] = Counter()
    with pdfplumber.open(BytesIO(raw)) as pdf:
        for page_index, page in enumerate(pdf.pages):
            page_text = page.extract_text() or ""
            texts.append(page_text)
            page_sizes.append((page.width, page.height))
            edge_text = "".join(
                char.get("text", "")
                for char in page.chars
                if char.get("top", 1000) < 55 or char.get("bottom", 0) > page.height - 45
            )
            if re.search(r"[0-9๐-๙]", edge_text):
                edge_number_pages += 1
            if page_index < 12 and re.search(r"ใบรับรองวิทยานิพนธ์|คณะกรรมการสอบวิทยานิพนธ์", page_text, re.I):
                approval_found = True
            for char in page.chars:
                try:
                    font_sizes[round(float(char.get("size", 0)), 1)] += 1
                except (TypeError, ValueError):
                    continue
        page_count = len(pdf.pages)
        full_text = "\n\n".join(texts)

    checks = []
    a4_pages = sum(1 for width, height in page_sizes if _near(width, 595.28, 5) and _near(height, 841.89, 5))
    checks.append(_check("กระดาษ A4", "pass" if a4_pages == len(page_sizes) else "review", f"พบ {a4_pages} จาก {len(page_sizes)} หน้าเป็นขนาด A4 โดยประมาณ"))
    checks.append(_check("ระยะขอบ", "review", "ประเมินตำแหน่งขอบข้อความจาก PDF ได้โดยประมาณเท่านั้น; ตรวจเพิ่มจากแต่ละหน้า โดยเฉพาะหน้าขึ้นบทใหม่"))
    if font_sizes:
        common_size, count = font_sizes.most_common(1)[0]
        checks.append(_check("ขนาดตัวอักษร", "pass" if min(abs(common_size - 16), abs(common_size - 12)) <= 1 else "review", f"ขนาดที่พบบ่อย {common_size:g} พอยต์ ({count} อักขระ); คู่มือระบุ 16 พอยต์ภาษาไทย และ 12 พอยต์ภาษาอังกฤษ"))
    else:
        checks.append(_check("ขนาดตัวอักษร", "review", "ไม่พบข้อมูลตัวอักษร; PDF อาจเป็นภาพสแกน"))
    number_ratio = edge_number_pages / max(1, page_count)
    checks.append(_check("ตำแหน่งเลขหน้า", "pass" if number_ratio >= 0.7 else "review", f"พบตัวเลขบริเวณขอบบน/ล่างใน {edge_number_pages} จาก {page_count} หน้า; ควรตรวจตำแหน่งและลำดับเลขหน้าด้วยตนเอง"))
    checks.append(_check("หน้าอนุมัติ/ใบรับรอง", "pass" if approval_found else "review", "พบข้อความที่บ่งชี้หน้าใบรับรองใน 12 หน้าแรก" if approval_found else "ไม่พบข้อความที่บ่งชี้หน้าใบรับรองใน 12 หน้าแรก หรือ PDF เป็นภาพสแกน"))
    return full_text, checks, {"page_count": page_count}


def analyze_document(filename: str, raw: bytes, citation_style: str = "auto") -> dict[str, Any]:
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError("ไฟล์มีขนาดเกิน 4.4 MB")
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "docx":
        text, format_checks, metadata = _analyze_docx(raw)
        file_type = "Word (.docx)"
    elif suffix == "pdf":
        text, format_checks, metadata = _analyze_pdf(raw)
        file_type = "PDF"
    else:
        raise ValueError("รองรับเฉพาะไฟล์ Word .docx และ PDF เท่านั้น")
    if not text.strip():
        raise ValueError("อ่านเนื้อหาในไฟล์ไม่พบ; หากเป็น PDF สแกนภาพให้แปลงเป็น PDF ที่เลือกข้อความได้ก่อน")
    return {
        "filename": filename,
        "file_type": file_type,
        "extracted_characters": len(text),
        "format_checks": format_checks,
        "citation_audit": _citation_audit(text, citation_style),
        "outline_audit": _outline_audit(text),
        "metadata": metadata,
        "text_for_ai": text,
    }


def build_request_docx(form_type: str, details: dict[str, str]) -> BytesIO:
    form_titles = {
        "advisor": "คำร้องขอแต่งตั้ง/เปลี่ยนแปลงอาจารย์ที่ปรึกษาวิทยานิพนธ์",
        "topic_exam": "คำร้องขอสอบหัวข้อวิทยานิพนธ์",
        "progress_exam": "คำร้องขอสอบความก้าวหน้าวิทยานิพนธ์",
        "defense_exam": "คำร้องขอสอบป้องกันวิทยานิพนธ์",
        "submit_thesis": "คำร้องขอส่งวิทยานิพนธ์ฉบับสมบูรณ์",
    }
    if form_type not in form_titles:
        raise ValueError("ไม่รู้จักประเภทคำร้องนี้")

    document = Document()
    document.add_heading("แบบร่างคำร้อง (ไม่ใช่แบบฟอร์มทางการ)", 0)
    document.add_paragraph("โปรดตรวจสอบและนำข้อมูลไปกรอกในแบบฟอร์มฉบับทางการของบัณฑิตวิทยาลัยก่อนยื่น")
    document.add_heading(form_titles[form_type], level=1)
    document.add_paragraph(f"วันที่: {details.get('request_date', '') or '................................'}")
    document.add_paragraph("เรียน คณบดีบัณฑิตวิทยาลัย")
    document.add_paragraph(
        f"ข้าพเจ้า {details.get('student_name', '') or '................................'} "
        f"รหัสนักศึกษา {details.get('student_id', '') or '................................'} "
        f"หลักสูตร/สาขาวิชา {details.get('program', '') or '................................'}"
    )
    document.add_paragraph(f"ชื่อวิทยานิพนธ์: {details.get('thesis_title', '') or '................................'}")
    document.add_paragraph(f"อาจารย์ที่ปรึกษา: {details.get('advisor', '') or '................................'}")
    document.add_paragraph(f"ความประสงค์/รายละเอียด: {details.get('request_detail', '') or '................................'}")
    document.add_paragraph("จึงเรียนมาเพื่อโปรดพิจารณา")
    document.add_paragraph("\nลงชื่อ ........................................................ ผู้ยื่นคำร้อง")
    document.add_paragraph(f"({details.get('student_name', '') or '................................'})")
    document.add_paragraph("\nความเห็นอาจารย์ที่ปรึกษา: ........................................................................")
    document.add_paragraph("\nผลการพิจารณา/สำหรับเจ้าหน้าที่: ................................................................")

    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return stream
